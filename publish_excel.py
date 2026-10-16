
from pathlib import Path
import csv
import subprocess
import sys
from openpyxl import load_workbook

# CHANGE THESE SETTINGS
EXCEL_FILE = Path(r"C:\ExcelData\MyData.xlsx")
SHEET_NAME = "Sheet1"
REPO_FOLDER = Path(r"C:\GitHub\ahmedabad-city-division")
CSV_FILE = REPO_FOLDER / "data.csv"


def run_git(*args):
    subprocess.run(
        ["git", *args],
        cwd=str(REPO_FOLDER),
        check=True
    )


def main():
    if not EXCEL_FILE.is_file():
        raise FileNotFoundError(f"Excel file not found: {EXCEL_FILE}")

    if not (REPO_FOLDER / ".git").exists():
        raise RuntimeError("Repository folder is not a Git repository.")

    # Read cached formula results where available.
    workbook = load_workbook(
        EXCEL_FILE,
        read_only=True,
        data_only=True
    )

    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(
                f"Worksheet '{SHEET_NAME}' not found. "
                f"Available sheets: {workbook.sheetnames}"
            )

        sheet = workbook[SHEET_NAME]

        with CSV_FILE.open(
            "w", newline="", encoding="utf-8-sig"
        ) as file:
            writer = csv.writer(file)

            for row in sheet.iter_rows(values_only=True):
                writer.writerow([
                    "" if value is None else value
                    for value in row
                ])
    finally:
        workbook.close()

    # Stage only the generated CSV.
    run_git("add", "data.csv")

    # Do not create an empty commit when data has not changed.
    result = subprocess.run(
        ["git", "diff", "--cached", "--quiet", "--", "data.csv"],
        cwd=str(REPO_FOLDER)
    )

    if result.returncode == 0:
        print("No data changes. Nothing to publish.")
        return

    if result.returncode != 1:
        raise RuntimeError("Unable to check staged CSV changes.")

    run_git("commit", "-m", "Update website data from Excel")
    run_git("push", "origin", "HEAD")

    print("Excel data published successfully.")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        sys.exit(1)
